from __future__ import annotations

from collections.abc import Generator
from io import BytesIO
import os
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlalchemy import create_engine
from sqlalchemy.orm import Session, sessionmaker

from backend.app.core.config import get_settings
from backend.app.db.models import Base
from backend.app.db.session import get_db
from backend.app.main import create_app


pytestmark = pytest.mark.live


@pytest.fixture(scope="module", autouse=True)
def require_live_configuration() -> None:
    if os.getenv("RUN_LIVE_TESTS") != "1":
        pytest.skip("Definí RUN_LIVE_TESTS=1 para habilitar llamadas reales acotadas.")
    settings = get_settings()
    missing = [
        name
        for name, value in (
            ("DEEPSEEK_API_KEY", settings.deepseek_api_key),
            ("TAVILY_API_KEY", settings.search_provider_api_key),
            ("GEMINI_API_KEY", settings.gemini_api_key),
        )
        if not value
    ]
    if missing:
        pytest.skip("Faltan claves para el smoke test: " + ", ".join(missing))


@pytest.fixture()
def live_client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Generator[TestClient, None, None]:
    database_path = tmp_path / "live-smoke.db"
    storage_path = tmp_path / "cvs"
    monkeypatch.setenv("DATABASE_URL", f"sqlite:///{database_path.as_posix()}")
    monkeypatch.setenv("CV_STORAGE_DIR", str(storage_path))
    monkeypatch.setenv("SEARCH_PROVIDER", "tavily")
    monkeypatch.setenv("AGENT_MAX_MODEL_TURNS", "6")
    monkeypatch.setenv("AGENT_MAX_SEARCHES", "1")
    monkeypatch.setenv("AGENT_MAX_INSPECTIONS", "2")
    monkeypatch.setenv("AGENT_MAX_ELAPSED_SECONDS", "120")

    engine = create_engine(
        f"sqlite:///{database_path.as_posix()}", connect_args={"check_same_thread": False}
    )
    Base.metadata.create_all(engine)
    testing_session = sessionmaker(bind=engine, autoflush=False, autocommit=False)

    def override_get_db() -> Generator[Session, None, None]:
        with testing_session() as db:
            yield db

    app = create_app()
    app.dependency_overrides[get_db] = override_get_db
    client = TestClient(app)
    yield client
    client.close()
    engine.dispose()


def create_conversation(client: TestClient, title: str) -> str:
    response = client.post("/api/conversations", json={"title": title})
    assert response.status_code == 201
    return response.json()["conversation_id"]


def synthetic_cv() -> bytes:
    document = Document()
    document.add_heading("Alex Rivera", 0)
    document.add_paragraph("Estudiante de ingeniería informática con experiencia académica en Python y SQL.")
    document.add_heading("Experiencia", level=1)
    document.add_paragraph("Proyecto universitario: API local con FastAPI, SQLite y pruebas automatizadas.")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_bounded_live_agent_research_and_cv_review(live_client: TestClient) -> None:
    direct_id = create_conversation(live_client, "Consejo")
    direct = live_client.post(
        f"/api/conversations/{direct_id}/messages",
        json={"content": "Dame un consejo breve y general para una entrevista junior.", "attachments": []},
    )
    assert direct.status_code == 202
    direct_detail = live_client.get(f"/api/conversations/{direct_id}").json()
    assert direct_detail["current_task"]["status"] == "completed"
    assert direct_detail["current_task"]["usage"]["provider_requests"] >= 1

    research_id = create_conversation(live_client, "Investigación")
    research = live_client.post(
        f"/api/conversations/{research_id}/messages",
        json={
            "content": (
                "Investigá brevemente a Mercado Libre como empleador para un perfil tecnológico "
                "junior. Usá exactamente una búsqueda web y como máximo dos inspecciones. "
                "Respondé solo con afirmaciones respaldadas por citas."
            ),
            "attachments": [],
        },
    )
    assert research.status_code == 202
    research_detail = live_client.get(f"/api/conversations/{research_id}").json()
    assert research_detail["current_task"]["status"] == "completed"
    assistant = [item for item in research_detail["messages"] if item["role"] == "assistant"][-1]
    assert assistant["citations"]
    assert research_detail["current_task"]["usage"]["searches"] == 1
    assert research_detail["current_task"]["usage"]["inspections"] <= 2

    uploaded = live_client.post(
        "/api/cvs",
        files={
            "file": (
                "synthetic-cv.docx",
                synthetic_cv(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
        data={"display_name": "CV sintético"},
    )
    assert uploaded.status_code == 201
    cv = uploaded.json()
    cv_conversation_id = create_conversation(live_client, "CV")
    prepared = live_client.post(
        f"/api/conversations/{cv_conversation_id}/messages",
        json={
            "content": "Prepará recomendaciones verificables para adaptar este CV al puesto adjunto.",
            "attachments": [
                {"type": "cv", "artifact_id": cv["cv_id"]},
                {
                    "type": "job_description",
                    "title": "Backend trainee",
                    "content": "Buscamos trainee con Python, SQL, APIs REST, trabajo en equipo e inglés.",
                },
            ],
        },
    )
    assert prepared.status_code == 202
    paused = live_client.get(f"/api/conversations/{cv_conversation_id}").json()
    assert paused["current_task"]["status"] == "awaiting_review"
    artifact_id = paused["current_task"]["pause"]["artifact_id"]
    artifact = live_client.get(f"/api/cv-recommendations/{artifact_id}").json()
    decisions = [
        {"suggestion_id": suggestion["id"], "decision": "rejected"}
        for suggestion in artifact["payload"].get("change_suggestions", [])
    ]
    reviewed = live_client.post(
        f"/api/task-runs/{prepared.json()['task_run_id']}/resume",
        json={
            "response_type": "review",
            "artifact_id": artifact_id,
            "review_decisions": decisions,
            "draft_text": artifact["draft_text"],
            "save_as_cv_version": True,
        },
    )
    assert reviewed.status_code == 202
    completed = live_client.get(f"/api/conversations/{cv_conversation_id}").json()
    assert completed["current_task"]["status"] == "completed"
    cv_detail = live_client.get(f"/api/cvs/{cv['cv_id']}").json()
    assert len(cv_detail["versions"]) == 2
